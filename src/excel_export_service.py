"""
Excel Export Service — Populates the Excel template with extraction results.

Reads field mappings dynamically from the 'Field Mapping' sheet in the template
and a JSON config file for JSON-key-to-Excel-field-name bridging.
No hardcoded cell references.
"""

import json
import logging
import os
from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import openpyxl

logger = logging.getLogger(__name__)

# Paths resolved relative to project root
PROJECT_ROOT = Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "templates" / "Notice_Ninja_Extraction_Tracker_updated1.xlsx"
DEFAULT_CONFIG_PATH = PROJECT_ROOT / "excel_field_config.json"

# Sheet names in the template
FIELD_MAPPING_SHEET = "Field Mapping"
NOTICE_FORM_SHEET = "Notice Form"
TRACKER_SHEET = "Tracker"


class ExcelExportService:
    """
    Service that fills the Excel template with extracted notice data.

    All cell references and column positions are resolved dynamically
    from the template's Field Mapping sheet and the Tracker header row.
    """

    def __init__(
        self,
        template_path: Optional[Path] = None,
        config_path: Optional[Path] = None,
    ) -> None:
        """
        Initialize the export service.

        Args:
            template_path: Path to the Excel template file.
            config_path: Path to the JSON field config file.
        """
        self.template_path = template_path or DEFAULT_TEMPLATE_PATH
        self.config_path = config_path or DEFAULT_CONFIG_PATH

        if not self.template_path.exists():
            raise FileNotFoundError(f"Excel template not found: {self.template_path}")
        if not self.config_path.exists():
            raise FileNotFoundError(f"Field config not found: {self.config_path}")

        self._field_config = self._load_field_config()
        logger.info(
            "ExcelExportService initialized — template: %s, config: %s",
            self.template_path,
            self.config_path,
        )

    # ------------------------------------------------------------------
    # Config loading
    # ------------------------------------------------------------------

    def _load_field_config(self) -> Dict[str, str]:
        """Load the JSON-key → Excel-field-name mapping from the config file."""
        with open(self.config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        logger.debug("Loaded field config with %d mappings", len(config))
        return config

    # ------------------------------------------------------------------
    # Template mapping readers
    # ------------------------------------------------------------------

    def _read_field_mapping(
        self, wb: openpyxl.Workbook
    ) -> Tuple[Dict[str, str], Dict[str, str]]:
        """
        Read the 'Field Mapping' sheet and return two lookup dicts.

        Returns:
            Tuple of:
                - form_cell_map: {field_name: cell_reference}  e.g. {"Account": "A6"}
                - tracker_col_map: {field_name: tracker_column_name}  e.g. {"Classification": "Classification"}
        """
        ws = wb[FIELD_MAPPING_SHEET]
        form_cell_map: Dict[str, str] = {}
        tracker_col_map: Dict[str, str] = {}

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            # Columns: A=Section, B=Field, C=Excel Form Cell, D=Tracker Column, E=Source
            field_name = row[1]
            form_cell = row[2]
            tracker_col = row[3]

            if not field_name:
                continue

            field_name = str(field_name).strip()

            if form_cell:
                form_cell_map[field_name] = str(form_cell).strip()

            if tracker_col and str(tracker_col).strip() != "Not in current Tracker":
                tracker_col_map[field_name] = str(tracker_col).strip()

        logger.debug(
            "Field Mapping loaded — %d form cells, %d tracker columns",
            len(form_cell_map),
            len(tracker_col_map),
        )
        return form_cell_map, tracker_col_map

    def _read_tracker_headers(self, wb: openpyxl.Workbook) -> Dict[str, int]:
        """
        Read the Tracker sheet row 1 headers and return a name→column-index map.

        Returns:
            {header_name: column_index}  e.g. {"Classification": 32, "Tax Period": 35}
        """
        ws = wb[TRACKER_SHEET]
        header_map: Dict[str, int] = {}

        for col_idx in range(1, ws.max_column + 1):
            header_val = ws.cell(row=1, column=col_idx).value
            if header_val:
                header_map[str(header_val).strip()] = col_idx

        logger.debug("Tracker headers loaded — %d columns", len(header_map))
        return header_map

    # ------------------------------------------------------------------
    # Sheet population
    # ------------------------------------------------------------------

    def _fill_notice_form(
        self,
        wb: openpyxl.Workbook,
        extracted_data: Dict[str, Any],
        form_cell_map: Dict[str, str],
    ) -> int:
        """
        Populate the Notice Form sheet with extracted data.

        Handles merged cells by unmerging the range and writing to the
        top-left cell of that range.

        Args:
            wb: The workbook to fill.
            extracted_data: The extraction JSON dict.
            form_cell_map: {field_name: cell_reference} from Field Mapping.

        Returns:
            Number of fields populated.
        """
        ws = wb[NOTICE_FORM_SHEET]
        populated = 0

        for json_key, value in extracted_data.items():
            if value is None:
                continue

            # Resolve JSON key → Excel field name
            excel_field_name = self._field_config.get(json_key)
            if not excel_field_name:
                logger.debug("No config mapping for JSON key: %s", json_key)
                continue

            # Resolve Excel field name → form cell reference
            cell_ref = form_cell_map.get(excel_field_name)
            if not cell_ref:
                logger.debug(
                    "No form cell mapping for field: %s (json key: %s)",
                    excel_field_name,
                    json_key,
                )
                continue

            # Handle merged cells: find and unmerge any range containing this cell,
            # then write to the top-left cell of that range
            write_cell = cell_ref
            from openpyxl.cell.cell import MergedCell

            if isinstance(ws[cell_ref], MergedCell):
                for merge_range in list(ws.merged_cells.ranges):
                    if cell_ref in merge_range:
                        ws.unmerge_cells(str(merge_range))
                        # Use the top-left cell of the former merge range
                        write_cell = merge_range.start_cell.coordinate
                        logger.debug(
                            "Unmerged %s for field %s, writing to %s",
                            merge_range,
                            excel_field_name,
                            write_cell,
                        )
                        break

            # Write value to the cell
            ws[write_cell] = value
            populated += 1
            logger.debug(
                "Notice Form: %s → cell %s = %s",
                excel_field_name,
                write_cell,
                value,
            )

        logger.info("Notice Form populated — %d fields written", populated)
        return populated

    def _fill_tracker(
        self,
        wb: openpyxl.Workbook,
        extracted_data: Dict[str, Any],
        tracker_col_map: Dict[str, str],
        header_map: Dict[str, int],
    ) -> int:
        """
        Populate the Tracker sheet row 2 with extracted data.

        Args:
            wb: The workbook to fill.
            extracted_data: The extraction JSON dict.
            tracker_col_map: {field_name: tracker_column_name} from Field Mapping.
            header_map: {header_name: column_index} from Tracker row 1.

        Returns:
            Number of fields populated.
        """
        ws = wb[TRACKER_SHEET]
        data_row = 2  # First data row (row 1 = headers)
        populated = 0

        for json_key, value in extracted_data.items():
            if value is None:
                continue

            # Resolve JSON key → Excel field name
            excel_field_name = self._field_config.get(json_key)
            if not excel_field_name:
                continue

            # Resolve Excel field name → tracker column name
            tracker_col_name = tracker_col_map.get(excel_field_name)
            if not tracker_col_name:
                continue

            # Resolve tracker column name → column index via headers
            col_idx = header_map.get(tracker_col_name)
            if not col_idx:
                logger.debug(
                    "Tracker header not found for column name: %s",
                    tracker_col_name,
                )
                continue

            # Write value to the data row
            ws.cell(row=data_row, column=col_idx, value=value)
            populated += 1
            logger.debug(
                "Tracker: %s → col %d (row %d) = %s",
                tracker_col_name,
                col_idx,
                data_row,
                value,
            )

        logger.info("Tracker populated — %d fields written to row %d", populated, data_row)
        return populated

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def generate(self, extracted_data: Dict[str, Any], file_name: str = "") -> bytes:
        """
        Generate a filled Excel workbook from the template.

        Args:
            extracted_data: The extracted notice data dict (18 fields).
            file_name: Original document filename (for logging).

        Returns:
            The filled workbook as bytes (.xlsx).
        """
        logger.info("Generating filled Excel for: %s", file_name or "(unknown)")

        # Load template workbook
        wb = openpyxl.load_workbook(self.template_path)

        # Read dynamic mappings
        form_cell_map, tracker_col_map = self._read_field_mapping(wb)
        header_map = self._read_tracker_headers(wb)

        # Populate both sheets
        form_count = self._fill_notice_form(wb, extracted_data, form_cell_map)
        tracker_count = self._fill_tracker(wb, extracted_data, tracker_col_map, header_map)

        logger.info(
            "Excel generation complete — Notice Form: %d fields, Tracker: %d fields",
            form_count,
            tracker_count,
        )

        # Remove helper sheets — output should contain only the Notice Form
        for sheet_name in [TRACKER_SHEET, FIELD_MAPPING_SHEET]:
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
                logger.debug("Removed sheet '%s' from output workbook", sheet_name)

        # Save to in-memory buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
