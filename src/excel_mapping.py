import os
import openpyxl

EXCEL_PATH = r"c:\Users\c1755\Notice-extraction\Notice_Mapping_Extracted_IT Team2.xlsx"

def load_mappings():
    """
    Reads the Notice Mapping Excel file and returns:
    - notice_type_to_manager: dict mapping Notice Type -> Notice Manager
    - allowed_notice_types: list of all valid Notice Types
    """
    notice_type_to_manager = {}
    allowed_notice_types = []
    
    if os.path.exists(EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
            if "Notice Mapping" in wb.sheetnames:
                ws = wb["Notice Mapping"]
                for row in ws.iter_rows(min_row=2):
                    classification = row[0].value
                    notice_type = row[1].value
                    notice_manager = row[2].value
                    
                    if notice_type:
                        notice_type_str = str(notice_type).strip()
                        allowed_notice_types.append(notice_type_str)
                        if notice_manager:
                            notice_type_to_manager[notice_type_str] = str(notice_manager).strip()
        except Exception as e:
            print(f"Warning: Failed to load Excel mapping: {e}")
            
    return notice_type_to_manager, list(set(allowed_notice_types))
